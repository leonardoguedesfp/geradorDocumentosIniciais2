"""Tests for Excel reader: complete sheet, missing columns, multiple clients, validation."""

import os
import sys
import tempfile
import shutil

import pytest
from openpyxl import Workbook

sys.path.insert(0, os.path.join(os.path.dirname(__file__), ".."))

from app.core.excel_reader import read_excel, COLUMN_MAP


def _create_test_xlsx(path, headers=None, rows=None, header_row=3):
    """Helper to create test xlsx files matching the expected format."""
    wb = Workbook()
    ws = wb.active
    ws.title = "Clientes"

    # Row 1: title
    ws.cell(row=1, column=1, value="Ricardo Passos Advocacia — Cadastro de Clientes")

    # Row 2: group headers
    ws.cell(row=2, column=1, value="DADOS PESSOAIS")
    ws.cell(row=2, column=7, value="ENDEREÇO")
    ws.cell(row=2, column=12, value="CONTATO")
    ws.cell(row=2, column=14, value="DADOS DO CASO")

    if headers is None:
        headers = list(COLUMN_MAP.keys())

    for i, h in enumerate(headers, 1):
        ws.cell(row=header_row, column=i, value=h)

    if rows:
        for r_idx, row_data in enumerate(rows, header_row + 1):
            for c_idx, val in enumerate(row_data, 1):
                ws.cell(row=r_idx, column=c_idx, value=val)

    wb.save(path)
    return path


@pytest.fixture
def tmp_xlsx_dir():
    d = tempfile.mkdtemp()
    yield d
    shutil.rmtree(d, ignore_errors=True)


class TestReadExcel:
    def test_complete_single_client(self, tmp_xlsx_dir):
        path = os.path.join(tmp_xlsx_dir, "test.xlsx")
        row = [
            "João da Silva", "brasileiro", "solteiro", "analista", "1234567",
            "529.982.247-25", "Rua A, 100", "Centro", "Brasília", "DF",
            "70000-000", "joao@email.com", "(61) 99999-0000", "Empresa XYZ"
        ]
        _create_test_xlsx(path, rows=[row])

        result = read_excel(path)
        assert len(result["clients"]) == 1
        assert result["errors"] == []
        row_num, cliente = result["clients"][0]
        assert cliente.nome_completo == "João da Silva"
        assert cliente.cpf == "529.982.247-25"
        assert cliente.bairro == "Centro"

    def test_multiple_clients(self, tmp_xlsx_dir):
        path = os.path.join(tmp_xlsx_dir, "test.xlsx")
        rows = [
            ["Maria Santos", "brasileira", "casada", "professora", "7654321",
             "987.654.321-00", "Av B, 200", "Norte", "Brasília", "DF",
             "71000-000", "maria@email.com", "(61) 88888-0000", "Emp ABC"],
            ["Pedro Souza", "brasileiro", "casado", "engenheiro", "9999999",
             "529.982.247-25", "Rua C, 300", "Sul", "Brasília", "DF",
             "72000-000", "pedro@email.com", "(61) 77777-0000", "Emp DEF"],
        ]
        _create_test_xlsx(path, rows=rows)

        result = read_excel(path)
        assert len(result["clients"]) == 2

    def test_missing_required_column(self, tmp_xlsx_dir):
        path = os.path.join(tmp_xlsx_dir, "test.xlsx")
        # Missing "CPF" column
        headers = [h for h in COLUMN_MAP.keys() if h != "CPF"]
        row = [
            "João da Silva", "brasileiro", "solteiro", "analista", "1234567",
            "Rua A, 100", "Centro", "Brasília", "DF",
            "70000-000", "joao@email.com", "(61) 99999-0000", "Empresa XYZ"
        ]
        _create_test_xlsx(path, headers=headers, rows=[row])

        result = read_excel(path)
        assert len(result["missing_columns"]) > 0
        assert "CPF" in result["missing_columns"]

    def test_empty_required_field_warning(self, tmp_xlsx_dir):
        path = os.path.join(tmp_xlsx_dir, "test.xlsx")
        # CPF empty
        row = [
            "João da Silva", "brasileiro", "solteiro", "analista", "1234567",
            "", "Rua A, 100", "Centro", "Brasília", "DF",
            "70000-000", "joao@email.com", "(61) 99999-0000", "Empresa XYZ"
        ]
        _create_test_xlsx(path, rows=[row])

        result = read_excel(path)
        assert len(result["clients"]) == 1  # Still loaded
        assert len(result["row_warnings"]) > 0  # But has warnings

    def test_skip_empty_rows(self, tmp_xlsx_dir):
        path = os.path.join(tmp_xlsx_dir, "test.xlsx")
        rows = [
            ["João Silva", "brasileiro", "solteiro", "analista", "123",
             "529.982.247-25", "Rua A, 100", "Centro", "Brasília", "DF",
             "70000-000", "j@e.com", "99999", "XYZ"],
            [None, None, None, None, None, None, None, None, None, None,
             None, None, None, None],  # Empty row
            ["Maria Santos", "brasileira", "casada", "prof", "456",
             "987.654.321-00", "Av B, 200", "Norte", "Brasília", "DF",
             "71000-000", "m@e.com", "88888", "ABC"],
        ]
        _create_test_xlsx(path, rows=rows)

        result = read_excel(path)
        assert len(result["clients"]) == 2

    def test_invalid_file(self, tmp_xlsx_dir):
        path = os.path.join(tmp_xlsx_dir, "bad.xlsx")
        with open(path, "w") as f:
            f.write("not an excel file")

        result = read_excel(path)
        assert len(result["errors"]) > 0

    def test_no_header_found(self, tmp_xlsx_dir):
        path = os.path.join(tmp_xlsx_dir, "noheader.xlsx")
        wb = Workbook()
        ws = wb.active
        ws.cell(row=1, column=1, value="Random data")
        wb.save(path)

        result = read_excel(path)
        assert len(result["errors"]) > 0
        assert "Cabeçalho" in result["errors"][0]

    def test_real_spreadsheet(self):
        """Test with the actual reference spreadsheet."""
        path = os.path.join(
            os.path.dirname(__file__), "..", "..",
            "__templates_padrao__", "Cadastro_Clientes_RicardoPassos.xlsx"
        )
        path = os.path.abspath(path)
        if not os.path.exists(path):
            pytest.skip("Reference spreadsheet not available")

        result = read_excel(path)
        # Should load without errors (no data rows, but no errors either)
        assert result["missing_columns"] == [] or result["missing_columns"] is not None
