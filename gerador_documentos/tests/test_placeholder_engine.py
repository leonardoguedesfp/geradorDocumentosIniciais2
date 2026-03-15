"""Tests for placeholder substitution engine."""

import os
import re
import sys
import tempfile
import shutil

import pytest
from docx import Document
from docx.shared import Pt, RGBColor

sys.path.insert(0, os.path.join(os.path.dirname(__file__), ".."))

from app.core.placeholder_engine import (
    data_extenso, fill_document, build_output_filename, _replace_in_paragraph,
)
from app.core.template_loader import load_template
from datetime import date


class TestDataExtenso:
    def test_specific_date(self):
        d = date(2025, 10, 13)
        assert data_extenso(d) == "13 de outubro de 2025"

    def test_january(self):
        d = date(2026, 1, 1)
        assert data_extenso(d) == "1 de janeiro de 2026"

    def test_december(self):
        d = date(2025, 12, 25)
        assert data_extenso(d) == "25 de dezembro de 2025"

    def test_default_is_today(self):
        result = data_extenso()
        assert str(date.today().year) in result


class TestFragmentedRuns:
    def test_placeholder_split_across_runs(self, tmp_dir):
        """Test that placeholders split across multiple runs are handled."""
        doc = Document()
        p = doc.add_paragraph()
        # Simulate fragmented runs: {{NOME_ in one run, COMPLETO}} in another
        run1 = p.add_run("Nome: {{NOME_")
        run1.bold = True
        run2 = p.add_run("COMPLETO}}")
        run2.bold = True

        path = os.path.join(tmp_dir, "frag_test.docx")
        doc.save(path)

        out_path = os.path.join(tmp_dir, "frag_out.docx")
        fill_document(path, {"NOME_COMPLETO": "João Silva"}, out_path)

        result_doc = Document(out_path)
        text = result_doc.paragraphs[0].text
        assert "João Silva" in text
        assert "{{" not in text
        assert "}}" not in text

    def test_preserves_formatting(self, tmp_dir):
        """Test that formatting of the first run is preserved."""
        doc = Document()
        p = doc.add_paragraph()
        run = p.add_run("{{NOME_COMPLETO}}")
        run.bold = True
        run.font.size = Pt(14)

        path = os.path.join(tmp_dir, "fmt_test.docx")
        doc.save(path)

        out_path = os.path.join(tmp_dir, "fmt_out.docx")
        fill_document(path, {"NOME_COMPLETO": "Maria Santos"}, out_path)

        result_doc = Document(out_path)
        result_run = result_doc.paragraphs[0].runs[0]
        assert result_run.bold is True
        assert result_run.font.size == Pt(14)
        assert "Maria Santos" in result_run.text


class TestFillDocument:
    def test_all_placeholders_replaced(self, templates_dir, sample_client_data, tmp_dir):
        """Test that no {{...}} remains after substitution."""
        template_path = os.path.join(templates_dir, "Procuracao_TEMPLATE.docx")
        if not os.path.exists(template_path):
            pytest.skip("Template file not available")

        out_path = os.path.join(tmp_dir, "output.docx")
        fill_document(template_path, sample_client_data, out_path)

        doc = Document(out_path)
        for p in doc.paragraphs:
            text = "".join(r.text for r in p.runs)
            assert "{{" not in text, f"Remaining placeholder in: {text}"
        for table in doc.tables:
            for row in table.rows:
                for cell in row.cells:
                    for p in cell.paragraphs:
                        text = "".join(r.text for r in p.runs)
                        assert "{{" not in text, f"Remaining placeholder in table: {text}"

    def test_missing_data_replaced_with_empty(self, tmp_dir):
        """Placeholder without data should be replaced with empty string."""
        doc = Document()
        doc.add_paragraph("{{NOME_COMPLETO}} and {{UNKNOWN_FIELD}}")
        path = os.path.join(tmp_dir, "missing_test.docx")
        doc.save(path)

        out_path = os.path.join(tmp_dir, "missing_out.docx")
        fill_document(path, {"NOME_COMPLETO": "Test"}, out_path)

        result = Document(out_path)
        text = result.paragraphs[0].text
        assert "Test" in text
        assert "{{" not in text

    def test_table_placeholders(self, tmp_dir):
        """Test placeholders in table cells are replaced."""
        doc = Document()
        table = doc.add_table(rows=1, cols=2)
        table.cell(0, 0).text = "{{NOME_COMPLETO}}"
        table.cell(0, 1).text = "{{CPF}}"
        path = os.path.join(tmp_dir, "table_test.docx")
        doc.save(path)

        out_path = os.path.join(tmp_dir, "table_out.docx")
        fill_document(path, {"NOME_COMPLETO": "Ana", "CPF": "123"}, out_path)

        result = Document(out_path)
        assert result.tables[0].cell(0, 0).text == "Ana"
        assert result.tables[0].cell(0, 1).text == "123"

    def test_real_templates(self, templates_dir, sample_client_data, tmp_dir):
        """Test all three real templates produce clean output."""
        for fname in [
            "Procuracao_TEMPLATE.docx",
            "Declaracao_Hipossuficiencia_TEMPLATE.docx",
            "Contrato_Prestacao_Servicos_TEMPLATE.docx",
        ]:
            template_path = os.path.join(templates_dir, fname)
            if not os.path.exists(template_path):
                pytest.skip(f"Template {fname} not available")

            out_path = os.path.join(tmp_dir, f"out_{fname}")
            fill_document(template_path, sample_client_data, out_path)

            doc = Document(out_path)
            for p in doc.paragraphs:
                text = "".join(r.text for r in p.runs)
                assert "{{" not in text, f"Remaining in {fname}: {text}"


class TestBuildOutputFilename:
    def test_basic(self):
        name = build_output_filename("Rinaldo da Silva Soares", "procuracao")
        assert name == "Rinaldo_da_Silva_Soares_Procuracao.docx"

    def test_declaracao(self):
        name = build_output_filename("Maria Santos", "declaracao")
        assert name == "Maria_Santos_Declaracao_Hipossuficiencia.docx"

    def test_contrato(self):
        name = build_output_filename("João Souza", "contrato")
        assert name == "Joao_Souza_Contrato.docx"

    def test_duplicate_v2(self):
        existing = {"Joao_Silva_Procuracao.docx"}
        name = build_output_filename("João Silva", "procuracao", existing)
        assert name == "Joao_Silva_Procuracao_v2.docx"

    def test_duplicate_v3(self):
        existing = {
            "Joao_Silva_Procuracao.docx",
            "Joao_Silva_Procuracao_v2.docx",
        }
        name = build_output_filename("João Silva", "procuracao", existing)
        assert name == "Joao_Silva_Procuracao_v3.docx"

    def test_accents_removed(self):
        name = build_output_filename("José André da Conceição", "procuracao")
        assert name == "Jose_Andre_da_Conceicao_Procuracao.docx"
