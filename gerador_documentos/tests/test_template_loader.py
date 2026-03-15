"""Tests for template loading and validation."""

import os
import sys
import shutil
import tempfile

import pytest
from docx import Document

sys.path.insert(0, os.path.join(os.path.dirname(__file__), ".."))

from app.core.template_loader import (
    load_template, load_default_templates, check_templates_folder,
    TEMPLATE_FILES, EXPECTED_PLACEHOLDERS,
)


class TestLoadTemplate:
    def test_load_valid_template(self, templates_dir):
        path = os.path.join(templates_dir, "Procuracao_TEMPLATE.docx")
        info = load_template(path, "procuracao")
        assert info.loaded is True
        assert info.error == ""
        assert info.document is not None

    def test_load_missing_file(self, tmp_dir):
        path = os.path.join(tmp_dir, "nonexistent.docx")
        info = load_template(path, "procuracao")
        assert info.loaded is False
        assert "não encontrado" in info.error

    def test_load_corrupted_file(self, tmp_dir):
        path = os.path.join(tmp_dir, "bad.docx")
        with open(path, "w") as f:
            f.write("not a docx file")
        info = load_template(path, "procuracao")
        assert info.loaded is False
        assert info.error != ""

    def test_custom_flag(self, templates_dir):
        path = os.path.join(templates_dir, "Procuracao_TEMPLATE.docx")
        info = load_template(path, "procuracao", is_custom=True)
        assert info.is_custom is True
        assert info.loaded is True

    def test_missing_placeholders_detected(self, tmp_dir):
        """Template without expected placeholders should report missing."""
        doc = Document()
        doc.add_paragraph("This template has no placeholders.")
        path = os.path.join(tmp_dir, "empty_template.docx")
        doc.save(path)

        info = load_template(path, "procuracao")
        assert info.loaded is True
        assert info.missing_placeholders is not None
        assert len(info.missing_placeholders) > 0

    def test_all_placeholders_present(self, tmp_dir):
        """Template with all expected placeholders should have no missing."""
        doc = Document()
        for ph in EXPECTED_PLACEHOLDERS["declaracao"]:
            doc.add_paragraph(f"{{{{{ph}}}}}")
        path = os.path.join(tmp_dir, "complete.docx")
        doc.save(path)

        info = load_template(path, "declaracao")
        assert info.loaded is True
        assert info.missing_placeholders is None


class TestLoadDefaultTemplates:
    def test_all_three_loaded(self, templates_dir):
        results = load_default_templates(templates_dir)
        assert len(results) == 3
        for doc_type, info in results.items():
            assert info.loaded is True, f"{doc_type} should be loaded"

    def test_one_missing(self, templates_dir):
        os.remove(os.path.join(templates_dir, "Contrato_Prestacao_Servicos_TEMPLATE.docx"))
        results = load_default_templates(templates_dir)
        assert results["procuracao"].loaded is True
        assert results["declaracao"].loaded is True
        assert results["contrato"].loaded is False

    def test_empty_folder(self, tmp_dir):
        results = load_default_templates(tmp_dir)
        for doc_type, info in results.items():
            assert info.loaded is False


class TestCheckTemplatesFolder:
    def test_all_present(self, templates_dir):
        status = check_templates_folder(templates_dir)
        assert all(status.values())

    def test_one_missing(self, templates_dir):
        os.remove(os.path.join(templates_dir, "Procuracao_TEMPLATE.docx"))
        status = check_templates_folder(templates_dir)
        assert status["procuracao"] is False
        assert status["declaracao"] is True
        assert status["contrato"] is True

    def test_nonexistent_folder(self):
        status = check_templates_folder("/nonexistent/path")
        assert all(v is False for v in status.values())
