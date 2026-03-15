"""Script to create test fixture files."""

import os
from openpyxl import Workbook

FIXTURES_DIR = os.path.join(os.path.dirname(__file__), "fixtures")
os.makedirs(FIXTURES_DIR, exist_ok=True)


def create_mock_xlsx():
    """Create clientes_mock.xlsx test fixture."""
    wb = Workbook()
    ws = wb.active
    ws.title = "Clientes"

    ws.cell(row=1, column=1, value="Ricardo Passos Advocacia — Cadastro de Clientes")

    ws.cell(row=2, column=1, value="DADOS PESSOAIS")
    ws.cell(row=2, column=7, value="ENDEREÇO")
    ws.cell(row=2, column=12, value="CONTATO")
    ws.cell(row=2, column=14, value="DADOS DO CASO")

    headers = [
        "Nome Completo", "Nacionalidade", "Estado Civil", "Profissão", "RG", "CPF",
        "Endereço (Logradouro e Número)", "Bairro", "Cidade", "UF", "CEP",
        "E-mail", "Telefone(s)", "Parte Contrária / Réu",
    ]
    for i, h in enumerate(headers, 1):
        ws.cell(row=3, column=i, value=h)

    clients = [
        ["João da Silva Souza", "brasileiro", "solteiro", "analista de sistemas",
         "1234567", "529.982.247-25", "Rua das Flores, 123", "Asa Sul",
         "Brasília", "DF", "70000-000", "joao@email.com", "(61) 99999-0000",
         "Empresa XYZ Ltda"],
        ["Maria Aparecida dos Santos", "brasileira", "casada", "professora",
         "7654321", "987.654.321-00", "Av. Brasil, 456", "Asa Norte",
         "Brasília", "DF", "71000-000", "maria@email.com", "(61) 88888-0000",
         "Empresa ABC S.A."],
        ["Pedro Henrique de Oliveira", "brasileiro", "divorciado", "engenheiro civil",
         "9999999", "052.004.068-02", "SQN 308 Bloco A, Apt 101", "Asa Norte",
         "Brasília", "DF", "72000-000", "pedro@email.com", "(61) 77777-0000",
         "Construtora DEF Ltda"],
    ]

    for r_idx, client in enumerate(clients, 4):
        for c_idx, val in enumerate(client, 1):
            ws.cell(row=r_idx, column=c_idx, value=val)

    path = os.path.join(FIXTURES_DIR, "clientes_mock.xlsx")
    wb.save(path)
    print(f"Created {path}")


if __name__ == "__main__":
    create_mock_xlsx()
