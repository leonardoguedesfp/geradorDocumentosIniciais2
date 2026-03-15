"""Load and validate .docx templates."""

from __future__ import annotations

import os
import re
from dataclasses import dataclass, field
from typing import Optional, Any

from docx import Document


TEMPLATE_FILES = {
    "procuracao": "Procuracao_TEMPLATE.docx",
    "declaracao": "Declaracao_Hipossuficiencia_TEMPLATE.docx",
    "contrato": "Contrato_Prestacao_Servicos_TEMPLATE.docx",
}

EXPECTED_PLACEHOLDERS = {
    "procuracao": {
        "NOME_COMPLETO", "NACIONALIDADE", "ESTADO_CIVIL", "RG", "CPF",
        "ENDERECO_COMPLETO", "PARTE_CONTRARIA", "DATA_EXTENSO",
        "PROFISSAO", "EMAIL", "TELEFONE",
    },
    "declaracao": {
        "NOME_COMPLETO", "NACIONALIDADE", "ESTADO_CIVIL", "RG", "CPF",
        "ENDERECO_COMPLETO", "PARTE_CONTRARIA", "DATA_EXTENSO",
    },
    "contrato": {
        "NOME_COMPLETO", "NACIONALIDADE", "ESTADO_CIVIL", "RG", "CPF",
        "ENDERECO_COMPLETO", "PARTE_CONTRARIA", "DATA_EXTENSO",
    },
}


@dataclass
class TemplateInfo:
    doc_type: str
    path: str
    document: Any = None
    loaded: bool = False
    error: str = ""
    missing_placeholders: Optional[list[str]] = None
    is_custom: bool = False


def _extract_placeholders_from_doc(doc: Document) -> set[str]:
    """Extract all {{PLACEHOLDER}} names from a Document."""
    placeholders = set()
    pattern = re.compile(r"\{\{(\w+)\}\}")

    for paragraph in doc.paragraphs:
        text = "".join(run.text for run in paragraph.runs)
        placeholders.update(pattern.findall(text))

    for table in doc.tables:
        for row in table.rows:
            for cell in row.cells:
                for paragraph in cell.paragraphs:
                    text = "".join(run.text for run in paragraph.runs)
                    placeholders.update(pattern.findall(text))

    return placeholders


def load_template(filepath: str, doc_type: str, is_custom: bool = False) -> TemplateInfo:
    """Load a single template file and validate its placeholders."""
    info = TemplateInfo(doc_type=doc_type, path=filepath, is_custom=is_custom)

    if not os.path.isfile(filepath):
        info.error = f"Arquivo não encontrado: {filepath}"
        return info

    try:
        doc = Document(filepath)
        info.document = doc
        info.loaded = True
    except Exception as e:
        info.error = f"Erro ao abrir template: {e}"
        return info

    found = _extract_placeholders_from_doc(doc)
    expected = EXPECTED_PLACEHOLDERS.get(doc_type, set())
    missing = expected - found
    if missing:
        info.missing_placeholders = sorted(missing)

    return info


def load_default_templates(folder: str) -> dict[str, TemplateInfo]:
    """Load all three templates from the default folder."""
    results = {}
    for doc_type, filename in TEMPLATE_FILES.items():
        filepath = os.path.join(folder, filename)
        results[doc_type] = load_template(filepath, doc_type)
    return results


def check_templates_folder(folder: str) -> dict[str, bool]:
    """Check which template files exist in the given folder."""
    status = {}
    for doc_type, filename in TEMPLATE_FILES.items():
        status[doc_type] = os.path.isfile(os.path.join(folder, filename))
    return status
